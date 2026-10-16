"""One-time end-to-end smoke test for PowerPoint -> private R2 slide images.

Generates its own 2-slide PPTX with an effectively impossible lesson ID,
converts via the actual LibreOffice/Poppler pipeline, validates both private
JPEGs, and removes only the exact temporary objects created by the check.
Never logs secrets or private keys.
"""

import hashlib
import io
import json
import os
import secrets
import sys

import boto3
from botocore.config import Config
from pptx import Presentation
from pptx.util import Inches

from app import ConversionRequest, convert_presentation


def main() -> int:
    bucket = os.environ["R2_BUCKET"]
    client = boto3.client(
        "s3",
        endpoint_url=os.environ["R2_ENDPOINT"],
        aws_access_key_id=os.environ["R2_ACCESS_KEY_ID"],
        aws_secret_access_key=os.environ["R2_SECRET_ACCESS_KEY"],
        region_name=os.environ.get("R2_REGION", "auto"),
        config=Config(connect_timeout=8, read_timeout=20, retries={"max_attempts": 2}),
    )

    lesson_id = 999999999999999999
    content_id = 999999999999999999
    random_file = secrets.token_hex(16)
    source_key = f"lessons/{lesson_id}/files/{random_file[:8]}-{random_file[8:12]}-4{random_file[13:16]}-8{random_file[17:20]}-{random_file[20:32]}.pptx"
    version = hashlib.sha256(source_key.encode("utf-8")).hexdigest()[:24]
    prefix = f"lessons/{lesson_id}/previews/{content_id}/{version}"
    expected_slides = [f"{prefix}/slides/slide-{i:04d}.jpg" for i in (1, 2)]
    expected_keys = [source_key, prefix + "/manifest.json", *expected_slides]

    success = False
    try:
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)

        for i in (1, 2):
            slide = presentation.slides.add_slide(presentation.slide_layouts[0])
            slide.shapes.title.text = f"YourSchool Test Slide {i}"
            slide.placeholders[1].text = "Temporary conversion test — not a student lesson."

        stream = io.BytesIO()
        presentation.save(stream)
        stream.seek(0)
        client.put_object(
            Bucket=bucket, Key=source_key, Body=stream.getvalue(),
            ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )

        req = ConversionRequest(
            lesson_id=lesson_id, content_id=content_id,
            source_key=source_key, preview_prefix=prefix,
        )
        convert_presentation(req)

        manifest_bytes = client.get_object(
            Bucket=bucket, Key=prefix + "/manifest.json",
        )["Body"].read()
        manifest = json.loads(manifest_bytes)
        if manifest.get("slides") != expected_slides:
            raise ValueError("Rendered slide manifest did not match 2 expected slides.")

        for slide_key in expected_slides:
            result = client.get_object(Bucket=bucket, Key=slide_key)
            if result["Body"].read(3)[:2] != b"\xff\xd8":
                raise ValueError("Converted slide is not JPEG.")

        success = True
    except Exception as exc:
        print("CONVERTER_PPTX_SMOKE_FAILED: " + type(exc).__name__, file=sys.stderr)
    finally:
        try:
            client.delete_objects(
                Bucket=bucket,
                Delete={"Objects": [{"Key": key} for key in expected_keys], "Quiet": True},
            )
        except Exception as exc:
            success = False
            print("CONVERTER_PPTX_CLEANUP_FAILED: " + type(exc).__name__, file=sys.stderr)

    if success:
        print("CONVERTER_PPTX_SMOKE_OK: 2 slides")
        return 0
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
